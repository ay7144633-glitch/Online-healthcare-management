from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

ROOT = Path(r"C:\claude works\Online-healthcare-management\docs")
SLIDES = ROOT / "rendered-slides"
PDF_PATH = ROOT / "presentation.pdf"
PPTX_PATH = ROOT / "presentation.pptx"
PAGE_WIDTH, PAGE_HEIGHT = 612, 792


def build_pdf() -> None:
    images = []
    contents = []
    for index in range(1, 15):
        image = Image.open(SLIDES / f"slide-{index:02d}.png").convert("RGB")
        buffer = BytesIO()
        image.save(buffer, format="JPEG", quality=95)
        data = buffer.getvalue()
        images.append((image.width, image.height, data))
        scale = min(PAGE_WIDTH / image.width, PAGE_HEIGHT / image.height)
        width, height = image.width * scale, image.height * scale
        x, y = (PAGE_WIDTH - width) / 2, (PAGE_HEIGHT - height) / 2
        contents.append(
            f"q\n{width:.2f} 0 0 {height:.2f} {x:.2f} {y:.2f} cm\n/Im{index + 2} Do\nQ\n".encode()
        )

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [] /Count 0 >>",
    ]
    for width, height, data in images:
        objects.append(
            f"<< /Type /XObject /Subtype /Image /Width {width} /Height {height} /ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /DCTDecode /Length {len(data)} >>\nstream\n".encode()
            + data
            + b"\nendstream"
        )

    page_references = []
    for index in range(14):
        image_number = index + 3
        content_number = index + 17
        page_references.append(len(objects) + 1)
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {PAGE_WIDTH} {PAGE_HEIGHT}] /Resources << /XObject << /Im{image_number} {image_number} 0 R >> >> /Contents {content_number} 0 R >>".encode()
        )
        objects.append(
            f"<< /Length {len(contents[index])} >>\nstream\n".encode()
            + contents[index]
            + b"endstream"
        )

    objects[1] = (
        f"<< /Type /Pages /Kids [{' '.join(f'{reference} 0 R' for reference in page_references)}] /Count {len(page_references)} >>".encode()
    )

    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode())
        output.extend(obj)
        output.extend(b"\nendobj\n")

    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    PDF_PATH.write_bytes(output)


def build_pptx() -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    for index in range(1, 15):
        slide = presentation.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        slide.shapes.add_picture(
            str(SLIDES / f"slide-{index:02d}.png"),
            Inches(0),
            Inches(0),
            Inches(13.333),
            Inches(7.5),
        )
    presentation.save(PPTX_PATH)


build_pdf()
build_pptx()
print(f"Created {PDF_PATH.name} and {PPTX_PATH.name}")
